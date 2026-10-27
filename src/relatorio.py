"""
Monta as saídas da semana:
  - relatorio.xlsx  (abas formatadas, pronto para quem quer mexer nos números)
  - relatorio.pdf   (2 páginas para ler/encaminhar)
  - RESUMO.md       (aparece direto no GitHub)
"""
from __future__ import annotations

from pathlib import Path

import matplotlib
import pandas as pd
from fpdf import FPDF
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

FONTE_DIR = Path(matplotlib.get_data_path()) / "fonts" / "ttf"  # DejaVu vem junto com o matplotlib
COR_CAB = "1F6FB2"


def _fmt_data(d: pd.Timestamp) -> str:
    return d.strftime("%d/%m/%Y")


# ------------------------------------------------------------------ Excel

def _formatar_aba(ws, formatos: dict[str, str]) -> None:
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=COR_CAB)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "A2"
    for i, col in enumerate(ws.iter_cols(min_row=1, max_row=1), start=1):
        nome = col[0].value
        largura = max(len(str(nome)), *(len(str(c.value)) for c in ws[get_column_letter(i)][1:50])) + 2
        ws.column_dimensions[get_column_letter(i)].width = min(largura, 28)
        if nome in formatos:
            for c in ws[get_column_letter(i)][1:]:
                c.number_format = formatos[nome]


def gerar_excel(arq: Path, abas: dict[str, pd.DataFrame]) -> Path:
    with pd.ExcelWriter(arq, engine="openpyxl") as w:
        for nome, df in abas.items():
            df.to_excel(w, sheet_name=nome, index=False)

    wb = load_workbook(arq)
    fmt = {
        "Preço atual": '"R$" #,##0.00', "Semana anterior": '"R$" #,##0.00', "Var. semana (R$)": '+0.00;-0.00;0.00',
        "Var. semana (%)": '+0.0"%";-0.0"%";0.0"%"', "Var. janela (%)": '+0.0"%";-0.0"%";0.0"%"',
        "Preço médio": '"R$" #,##0.00', "Mínimo": '"R$" #,##0.00', "Máximo": '"R$" #,##0.00',
        "Etanol": '"R$" #,##0.00', "Gasolina": '"R$" #,##0.00', "Paridade": "0.0%",
        "Semana": "dd/mm/yyyy",
    }
    for ws in wb.worksheets:
        _formatar_aba(ws, fmt)
    wb.save(arq)
    return arq


# ------------------------------------------------------------------ PDF

class _PDF(FPDF):
    def __init__(self, semana: str):
        super().__init__(format="A4")
        self.semana = semana
        self.add_font("DejaVu", "", str(FONTE_DIR / "DejaVuSans.ttf"))
        self.add_font("DejaVu", "B", str(FONTE_DIR / "DejaVuSans-Bold.ttf"))
        self.set_auto_page_break(True, margin=15)

    def header(self):
        self.set_font("DejaVu", "B", 9)
        self.set_text_color(120)
        self.cell(0, 6, f"Relatório semanal de preços de combustíveis · semana de {self.semana}", align="L")
        self.ln(9)
        self.set_text_color(0)

    def footer(self):
        self.set_y(-12)
        self.set_font("DejaVu", "", 7.5)
        self.set_text_color(130)
        self.cell(0, 6, "Fonte: ANP, Levantamento de Preços de Combustíveis. Gerado automaticamente.", align="L")
        self.cell(0, 6, f"{self.page_no()}", align="R")


def gerar_pdf(arq: Path, semana: pd.Timestamp, texto_resumo: str, origem: str,
              nacional: pd.DataFrame, figuras: dict[str, Path]) -> Path:
    pdf = _PDF(_fmt_data(semana))
    pdf.add_page()

    pdf.set_font("DejaVu", "B", 17)
    pdf.cell(0, 10, "Preços de combustíveis no Brasil", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("DejaVu", "", 10)
    pdf.set_text_color(90)
    pdf.cell(0, 6, f"Semana de {_fmt_data(semana)} a {_fmt_data(semana + pd.Timedelta(days=6))}",
             new_x="LMARGIN", new_y="NEXT")
    pdf.set_text_color(0)
    pdf.ln(3)

    pdf.set_font("DejaVu", "B", 11)
    pdf.cell(0, 7, "Resumo da semana" + (" (escrito por IA a partir dos números abaixo)" if origem == "ia" else ""),
             new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("DejaVu", "", 10)
    pdf.multi_cell(0, 5.4, texto_resumo)
    pdf.ln(3)

    # tabela nacional
    cab = ["Produto", "Preço atual", "Var. semana", f"Var. {int(nacional.semanas_janela.iloc[0])} sem.", "Mín – Máx"]
    larg = [55, 32, 28, 28, 47]
    pdf.set_font("DejaVu", "B", 9)
    pdf.set_fill_color(31, 111, 178)
    pdf.set_text_color(255)
    for c, w in zip(cab, larg):
        pdf.cell(w, 7, c, border=0, fill=True, align="C")
    pdf.ln()
    pdf.set_text_color(0)
    pdf.set_font("DejaVu", "", 9)
    for i, r in nacional.iterrows():
        pdf.set_fill_color(242, 245, 249) if i % 2 == 0 else pdf.set_fill_color(255, 255, 255)
        un = r.unidade.replace("R$", "")
        vals = [r.produto, f"R$ {r.preco_atual:.2f}{un}", f"{r.var_semana_pct:+.1f}%", f"{r.var_janela_pct:+.1f}%",
                f"R$ {r.preco_min:.2f} – {r.preco_max:.2f}"]
        for v, w, al in zip(vals, larg, "LRRRC"):
            pdf.cell(w, 6.5, v.replace(".", ","), fill=True, align=al)
        pdf.ln()
    pdf.ln(4)
    pdf.image(str(figuras["tendencia"]), w=185)

    # página 2: os dois gráficos por estado, um embaixo do outro
    pdf.add_page()
    pdf.image(str(figuras["gasolina"]), x=42, w=126)
    pdf.ln(2)
    pdf.image(str(figuras["paridade"]), x=42, w=126)
    pdf.output(str(arq))
    return arq


# ------------------------------------------------------------------ Markdown

def gerar_markdown(arq: Path, semana: pd.Timestamp, texto: str, origem: str, nacional: pd.DataFrame) -> Path:
    linhas = [
        f"# Combustíveis: semana de {_fmt_data(semana)}",
        "",
        f"> {texto}",
        "",
        f"<sub>Resumo {'escrito por IA a partir dos indicadores' if origem == 'ia' else 'gerado automaticamente a partir dos indicadores'}.</sub>",
        "",
        f"| Produto | Preço médio | Var. semana | Var. {int(nacional.semanas_janela.iloc[0])} semanas |",
        "|---|---:|---:|---:|",
    ]
    for r in nacional.itertuples():
        un = r.unidade.replace("R$", "")
        linhas.append(f"| {r.produto} | R$ {r.preco_atual:.2f}{un} | {r.var_semana_pct:+.1f}% | {r.var_janela_pct:+.1f}% |"
                      .replace(".", ","))
    linhas += ["", "![Tendência](tendencia_nacional.png)", "",
               "![Gasolina por estado](gasolina_por_estado.png)", "", "![Paridade](paridade_etanol.png)", ""]
    arq.write_text("\n".join(linhas), encoding="utf-8")
    return arq
